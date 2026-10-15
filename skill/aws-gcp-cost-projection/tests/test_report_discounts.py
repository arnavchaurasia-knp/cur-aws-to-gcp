import os
import csv
import duckdb
import pytest
from export_report import _QUERY_DISCOUNTS


@pytest.fixture
def mock_db_with_discounts(tmp_path):
    db_file = tmp_path / "projection.duckdb"
    conn = duckdb.connect(str(db_file))
    conn.execute("""
        CREATE TABLE aws_li_catalog (
            aws_li_key TEXT PRIMARY KEY,
            product TEXT,
            instance_type TEXT,
            aws_region TEXT,
            gcp_region TEXT,
            line_item_type TEXT,
            mechanic_group TEXT,
            pricing_model TEXT,
            is_workload BOOLEAN,
            total_usage DOUBLE,
            aws_amortized_cost DOUBLE
        );
        CREATE TABLE aws_li_to_gcp_li (
            aws_li_key TEXT,
            strategy TEXT,
            gcp_service TEXT,
            gcp_sku_id TEXT,
            component TEXT,
            unit_multiplier DOUBLE,
            projection_note TEXT,
            mapping_confidence DOUBLE
        );
        CREATE TABLE gcp_projection AS
        SELECT
            'k1' AS aws_li_key,
            'default' AS component,
            'Compute Engine' AS gcp_service,
            'n2-standard-4' AS gcp_machine_type,
            'us-central1' AS gcp_region,
            80.0 AS gcp_projected_cost,
            70.0 AS gcp_cost_1yr_cud,
            60.0 AS gcp_cost_3yr_cud,
            'mapped' AS projection_status,
            'OK' AS projection_note;
    """)

    conn.execute("""
        INSERT INTO aws_li_catalog VALUES
            ('k1', 'Amazon Elastic Compute Cloud', 'm5.xlarge', 'us-east-1', 'us-central1', 'Usage', 'compute_breakdown', 'OnDemand', true, 100, 100.0),
            ('k2', 'Amazon Elastic Compute Cloud', '', 'us-east-1', 'us-central1', 'SavingsPlanNegation', 'commitment_discount', 'SavingsPlan', false, 0, -25.0),
            ('k3', 'Amazon Relational Database Service', '', 'us-east-1', 'us-central1', 'Credit', 'negative_cost', 'Credit', false, 0, -10.0);
    """)
    conn.close()
    return db_file


def test_query_discounts_aggregation(mock_db_with_discounts):
    conn = duckdb.connect(str(mock_db_with_discounts), read_only=True)
    res = conn.execute(_QUERY_DISCOUNTS).fetchall()
    conn.close()

    # Should have 2 services with discounts
    assert len(res) == 2
    # EC2 discount: $25.0
    ec2_disc = next(r for r in res if "Compute Cloud" in r[1])
    assert ec2_disc[4] == 25.0
    # RDS discount: $10.0
    rds_disc = next(r for r in res if "Relational Database" in r[1])
    assert rds_disc[4] == 10.0


def test_discounts_in_report_render_html(tmp_path, monkeypatch):
    from render_report import main as render_main

    job_dir = tmp_path / "job"
    audit_dir = job_dir / "projection-audit"
    audit_dir.mkdir(parents=True)
    db_file = audit_dir / "projection.duckdb"

    conn = duckdb.connect(str(db_file))
    conn.execute("""
        CREATE TABLE aws_li_catalog (
            aws_li_key TEXT PRIMARY KEY,
            product TEXT,
            instance_type TEXT,
            instance_vcpus INTEGER,
            instance_ram_gb DOUBLE,
            instance_count DOUBLE,
            usage_type TEXT,
            operation TEXT,
            aws_region TEXT,
            gcp_region TEXT,
            line_item_type TEXT,
            mechanic_group TEXT,
            pricing_model TEXT,
            pricing_unit TEXT,
            is_workload BOOLEAN,
            total_usage DOUBLE,
            aws_amortized_cost DOUBLE
        );
        CREATE TABLE aws_li_to_gcp_li (
            aws_li_key TEXT,
            rowid BIGINT,
            strategy TEXT,
            gcp_service TEXT,
            gcp_machine_type TEXT,
            gcp_sku_id TEXT,
            gcp_sku_name TEXT,
            component TEXT,
            unit_multiplier DOUBLE,
            projection_note TEXT,
            mapping_confidence DOUBLE
        );
        CREATE TABLE gcp_projection (
            aws_li_key TEXT,
            component TEXT,
            gcp_service TEXT,
            gcp_machine_type TEXT,
            gcp_region TEXT,
            is_workload BOOLEAN,
            gcp_projected_cost DOUBLE,
            gcp_cost_1yr_cud DOUBLE,
            gcp_cost_3yr_cud DOUBLE,
            projection_status TEXT,
            projection_note TEXT
        );
        CREATE TABLE run_results (
            run_id TEXT,
            ts_utc TEXT,
            prospect TEXT
        );
    """)

    conn.execute("""
        INSERT INTO aws_li_catalog VALUES
            ('k1', 'Amazon Elastic Compute Cloud', 'm5.xlarge', 4, 16.0, 1.0, 'BoxUsage:m5.xlarge', 'RunInstances', 'us-east-1', 'us-central1', 'Usage', 'compute_breakdown', 'OnDemand', 'Hrs', true, 100, 1000.0),
            ('k2', 'Amazon Elastic Compute Cloud', '', NULL, NULL, NULL, 'SavingsPlanNegation', 'SP', 'us-east-1', 'us-central1', 'SavingsPlanNegation', 'commitment_discount', 'SavingsPlan', 'Hrs', false, 0, -200.0);
        INSERT INTO aws_li_to_gcp_li VALUES
            ('k1', 1, 'map', 'Compute Engine', 'n2-standard-4', 'sku-1', 'N2 Instance Core', 'default', 1.0, 'OK', 0.95);
        INSERT INTO gcp_projection VALUES
            ('k1', 'default', 'Compute Engine', 'n2-standard-4', 'us-central1', true, 850.0, 700.0, 550.0, 'mapped', 'OK');
        INSERT INTO run_results VALUES
            ('run-test', '2026-10-14 00:00:00', 'Acme Corp');
    """)
    conn.close()

    monkeypatch.chdir(job_dir)
    monkeypatch.setattr("render_report.JOB_DIR", str(job_dir))
    monkeypatch.setattr("render_report.DB_PATH", str(db_file))

    # Run render_report
    render_main()

    import glob
    html_files = glob.glob(str(audit_dir / "report-*.html"))
    assert len(html_files) > 0
    with open(html_files[0], encoding="utf-8") as f:
        html_content = f.read()

    # (a) Headline list-vs-list comparison
    assert "GRAND TOTAL (LIST PRICES)" in html_content
    assert "$1,000.00" in html_content  # AWS gross / list
    assert "$850.00" in html_content    # GCP OD

    # (b) Discounts by Service section
    assert "Discounts by Service" in html_content
    assert "−$200.00" in html_content

    # (c) Net AWS total context line
    assert "Customer Net AWS Total" in html_content
    assert "$800.00 / mo" in html_content


def test_export_report_discounts(tmp_path, monkeypatch):
    from export_report import main as export_main

    job_dir = tmp_path / "job_export"
    audit_dir = job_dir / "projection-audit"
    audit_dir.mkdir(parents=True)
    db_file = audit_dir / "projection.duckdb"

    conn = duckdb.connect(str(db_file))
    conn.execute("""
        CREATE TABLE aws_li_catalog (
            aws_li_key TEXT PRIMARY KEY,
            product TEXT,
            instance_type TEXT,
            usage_type TEXT,
            operation TEXT,
            aws_region TEXT,
            gcp_region TEXT,
            line_item_type TEXT,
            mechanic_group TEXT,
            pricing_model TEXT,
            pricing_unit TEXT,
            is_workload BOOLEAN,
            total_usage DOUBLE,
            aws_amortized_cost DOUBLE
        );
        CREATE TABLE aws_li_to_gcp_li (
            aws_li_key TEXT,
            rowid BIGINT,
            strategy TEXT,
            gcp_service TEXT,
            gcp_machine_type TEXT,
            gcp_sku_id TEXT,
            component TEXT,
            unit_multiplier DOUBLE,
            projection_note TEXT,
            mapping_confidence DOUBLE
        );
        CREATE TABLE gcp_projection (
            aws_li_key TEXT,
            component TEXT,
            gcp_service TEXT,
            gcp_machine_type TEXT,
            gcp_region TEXT,
            gcp_projected_cost DOUBLE,
            gcp_cost_1yr_cud DOUBLE,
            gcp_cost_3yr_cud DOUBLE,
            projection_status TEXT,
            projection_note TEXT
        );
        CREATE TABLE run_results (
            run_id TEXT,
            ts_utc TEXT,
            prospect TEXT
        );
    """)

    conn.execute("""
        INSERT INTO aws_li_catalog VALUES
            ('k1', 'Amazon Elastic Compute Cloud', 'm5.xlarge', 'BoxUsage', 'RunInstances', 'us-east-1', 'us-central1', 'Usage', 'compute_breakdown', 'OnDemand', 'Hrs', true, 100, 1000.0),
            ('k2', 'Amazon Elastic Compute Cloud', '', 'SavingsPlanNegation', 'SP', 'us-east-1', 'us-central1', 'SavingsPlanNegation', 'commitment_discount', 'SavingsPlan', 'Hrs', false, 0, -200.0);
        INSERT INTO aws_li_to_gcp_li VALUES
            ('k1', 1, 'map', 'Compute Engine', 'n2-standard-4', 'sku-1', 'default', 1.0, 'OK', 0.95);
        INSERT INTO gcp_projection VALUES
            ('k1', 'default', 'Compute Engine', 'n2-standard-4', 'us-central1', 850.0, 700.0, 550.0, 'mapped', 'OK');
        INSERT INTO run_results VALUES
            ('run-test', '2026-10-14 00:00:00', 'Acme Corp');
    """)
    conn.close()

    monkeypatch.chdir(job_dir)
    monkeypatch.setattr("export_report.DB_PATH", str(db_file))

    # Test CSV with --by-discounts
    csv_out = str(job_dir / "discounts.csv")
    monkeypatch.setattr("sys.argv", ["export_report.py", "--format", "csv", "--out", csv_out, "--by-discounts"])
    export_main()

    with open(csv_out, encoding="utf-8") as f:
        reader = list(csv.reader(f))
    assert len(reader) == 2  # header + 1 discount row
    assert reader[0] == ["#", "AWS Service", "Discount Type", "Line Items", "Discount Amount ($)"]
    assert reader[1][1] == "Amazon Elastic Compute Cloud"
    assert reader[1][4] == "200.0"

    # Test XLSX with Discounts by Service sheet
    import openpyxl
    xlsx_out = str(job_dir / "report.xlsx")
    monkeypatch.setattr("sys.argv", ["export_report.py", "--format", "xlsx", "--out", xlsx_out])
    export_main()

    wb = openpyxl.load_workbook(xlsx_out)
    assert "Discounts by Service" in wb.sheetnames
    ws = wb["Discounts by Service"]
    assert ws.cell(row=1, column=2).value == "AWS Service"
    assert ws.cell(row=2, column=2).value == "Amazon Elastic Compute Cloud"

